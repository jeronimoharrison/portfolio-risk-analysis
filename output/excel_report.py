"""Excel report generation with styled sheets."""

import pandas as pd
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers
from openpyxl.utils.dataframe import dataframe_to_rows
from pathlib import Path

from risk.returns import simple_returns, portfolio_returns, cumulative_returns
from risk.var import compute_all_var, per_asset_var
from risk.cvar import compute_all_cvar, per_asset_cvar
from risk.metrics import (
    annualized_return, annualized_volatility, sharpe_ratio, sortino_ratio,
    max_drawdown, compute_asset_metrics, compute_asset_class_metrics,
    risk_contribution,
)
from risk.correlation import correlation_matrix

# Styles
DARK_HEADER_FILL = PatternFill(start_color="1B2A3D", end_color="1B2A3D", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(bold=True, color="1B2A3D", size=14)
THIN_BORDER = Border(
    left=Side(style="thin", color="CCCCCC"),
    right=Side(style="thin", color="CCCCCC"),
    top=Side(style="thin", color="CCCCCC"),
    bottom=Side(style="thin", color="CCCCCC"),
)
GREEN_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
RED_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
YELLOW_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")

PCT_FORMAT = "0.00%"
NUM_FORMAT = "0.0000"
NUM2_FORMAT = "0.00"


def _style_header_row(ws, row_num: int, num_cols: int):
    """Apply dark header styling to a row."""
    for col in range(1, num_cols + 1):
        cell = ws.cell(row=row_num, column=col)
        cell.font = HEADER_FONT
        cell.fill = DARK_HEADER_FILL
        cell.alignment = Alignment(horizontal="center")
        cell.border = THIN_BORDER


def _write_df(ws, df: pd.DataFrame, start_row: int = 1, include_index: bool = True):
    """Write a DataFrame to a worksheet with styling."""
    rows = list(dataframe_to_rows(df, index=include_index, header=True))

    for r_idx, row in enumerate(rows, start_row):
        for c_idx, value in enumerate(row, 1):
            if value is None:
                continue
            cell = ws.cell(row=r_idx, column=c_idx)
            if isinstance(value, float):
                cell.value = value
                cell.number_format = NUM_FORMAT
            else:
                cell.value = value
            cell.border = THIN_BORDER

    # Style header row(s)
    _style_header_row(ws, start_row, len(rows[0]) if rows else 0)
    if include_index and len(rows) > 1:
        _style_header_row(ws, start_row + 1, len(rows[1]) if len(rows) > 1 else 0)

    return start_row + len(rows)


def _autofit_columns(ws):
    """Auto-fit column widths."""
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_len + 4, 30)


def _apply_conditional_colors(ws, col_idx: int, start_row: int, end_row: int):
    """Apply green/red fill based on positive/negative values."""
    for row in range(start_row, end_row + 1):
        cell = ws.cell(row=row, column=col_idx)
        if isinstance(cell.value, (int, float)):
            if cell.value > 0:
                cell.fill = GREEN_FILL
            elif cell.value < 0:
                cell.fill = RED_FILL


def _correlation_heatmap_fill(ws, start_row: int, start_col: int, n: int):
    """Apply heatmap coloring to correlation matrix cells."""
    for r in range(start_row, start_row + n):
        for c in range(start_col, start_col + n):
            cell = ws.cell(row=r, column=c)
            if isinstance(cell.value, (int, float)):
                val = cell.value
                if val >= 0.7:
                    cell.fill = PatternFill(start_color="63BE7B", end_color="63BE7B", fill_type="solid")
                elif val >= 0.3:
                    cell.fill = PatternFill(start_color="FFEB84", end_color="FFEB84", fill_type="solid")
                elif val >= -0.3:
                    cell.fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
                elif val >= -0.7:
                    cell.fill = PatternFill(start_color="F8B984", end_color="F8B984", fill_type="solid")
                else:
                    cell.fill = PatternFill(start_color="F8696B", end_color="F8696B", fill_type="solid")


def generate_report(
    portfolio: pd.DataFrame,
    asset_returns: pd.DataFrame,
    benchmark_returns: pd.DataFrame,
    weights: pd.Series,
    output_path: str = "portfolio_risk_report.xlsx",
) -> str:
    """Generate the full styled Excel risk report.

    Parameters
    ----------
    portfolio : pd.DataFrame
        Portfolio holdings.
    asset_returns : pd.DataFrame
        Daily asset returns.
    benchmark_returns : pd.DataFrame
        Daily benchmark returns.
    weights : pd.Series
        Portfolio weights indexed by ticker.
    output_path : str
        Output file path.

    Returns
    -------
    str
        Path to generated report.
    """
    from risk.returns import portfolio_returns as calc_port_returns

    port_ret = calc_port_returns(asset_returns, weights)

    wb = Workbook()

    # === Sheet 1: Summary ===
    ws_summary = wb.active
    ws_summary.title = "Summary"

    ws_summary.cell(row=1, column=1, value="Portfolio Risk Analysis Report").font = TITLE_FONT
    ws_summary.cell(row=2, column=1, value=f"Period: {asset_returns.index.min().date()} to {asset_returns.index.max().date()}")
    ws_summary.cell(row=3, column=1, value=f"Holdings: {len(portfolio)}")

    # Aggregate metrics
    metrics = {
        "Total Return": cumulative_returns(port_ret).iloc[-1] if len(port_ret) > 0 else 0,
        "Annualized Return": annualized_return(port_ret),
        "Annualized Volatility": annualized_volatility(port_ret),
        "Sharpe Ratio": sharpe_ratio(port_ret),
        "Sortino Ratio": sortino_ratio(port_ret),
        "Max Drawdown": max_drawdown(port_ret),
    }

    row = 5
    ws_summary.cell(row=row, column=1, value="Metric").font = HEADER_FONT
    ws_summary.cell(row=row, column=1).fill = DARK_HEADER_FILL
    ws_summary.cell(row=row, column=2, value="Value").font = HEADER_FONT
    ws_summary.cell(row=row, column=2).fill = DARK_HEADER_FILL
    row += 1

    for metric_name, value in metrics.items():
        ws_summary.cell(row=row, column=1, value=metric_name).border = THIN_BORDER
        cell = ws_summary.cell(row=row, column=2, value=value)
        cell.border = THIN_BORDER
        cell.number_format = PCT_FORMAT if "Return" in metric_name or "Volatility" in metric_name or "Drawdown" in metric_name else NUM2_FORMAT
        row += 1

    _autofit_columns(ws_summary)

    # === Sheet 2: Asset Class Breakdown ===
    ws_ac = wb.create_sheet("Asset Class Breakdown")
    ac_metrics = compute_asset_class_metrics(asset_returns, portfolio, weights)
    _write_df(ws_ac, ac_metrics, start_row=1, include_index=False)
    _autofit_columns(ws_ac)

    # === Sheet 3: VaR / CVaR ===
    ws_var = wb.create_sheet("VaR Analysis")

    ws_var.cell(row=1, column=1, value="Portfolio-Level VaR & CVaR").font = TITLE_FONT

    var_df = compute_all_var(port_ret, asset_returns, weights)
    next_row = _write_df(ws_var, var_df, start_row=3)

    cvar_df = compute_all_cvar(port_ret)
    next_row = _write_df(ws_var, cvar_df, start_row=next_row + 2)

    ws_var.cell(row=next_row + 1, column=1, value="Per-Asset VaR (95%)").font = TITLE_FONT
    asset_var = per_asset_var(asset_returns, 0.95)
    asset_cvar = per_asset_cvar(asset_returns, 0.95)
    combined = asset_var.join(asset_cvar)
    _write_df(ws_var, combined, start_row=next_row + 2)
    _autofit_columns(ws_var)

    # === Sheet 4: Correlation ===
    ws_corr = wb.create_sheet("Correlation")
    ws_corr.cell(row=1, column=1, value="Asset Correlation Matrix").font = TITLE_FONT

    corr = correlation_matrix(asset_returns)
    data_start = _write_df(ws_corr, corr, start_row=3)

    # Apply heatmap
    n = len(corr)
    _correlation_heatmap_fill(ws_corr, 5, 2, n)  # data starts at row 5, col 2
    _autofit_columns(ws_corr)

    # === Sheet 5: Holdings Detail ===
    ws_detail = wb.create_sheet("Holdings Detail")
    detail = compute_asset_metrics(asset_returns, benchmark_returns, portfolio)
    _write_df(ws_detail, detail, start_row=1, include_index=False)
    _autofit_columns(ws_detail)

    # Save
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path
